"""Excel export generator reading directly from PostgreSQL TKBIAuditEntry records."""

import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from app.db.models import AuditRun, TKBIAuditEntry


def generate_stamped_excel_buffer(audit_run: AuditRun, entries: list[TKBIAuditEntry]) -> io.BytesIO:
    """Generate Excel buffer formatted according to Template_Audit_TKBI.xlsx with live auditor edits."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    headers = [
        "Kode Emiten", "Sektor", "Bab", "KBLI", "TSCID", "TSC",
        "BentukJawaban", "Jawaban AI", "Keyakinan AI", "Reasoning AI",
        "Bukti (Sumber Dokumen)", "Auditor Feedback"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    hijau_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    transisi_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    tidak_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")

    for row_idx, item in enumerate(entries, start=2):
        effective_ans = item.auditor_override or item.jawaban_ai
        feedback_val = item.auditor_feedback
        if item.is_overridden:
            feedback_val = f"[OVERRIDE: {effective_ans}] {feedback_val or ''}"

        row_vals = [
            item.kode_emiten,
            item.sektor,
            item.bab,
            item.kbli,
            item.tsc_id,
            item.tsc,
            item.bentuk_jawaban,
            effective_ans,
            item.keyakinan_ai,
            item.reasoning_ai,
            item.bukti,
            feedback_val,
        ]
        ws.append(row_vals)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border_thin
            cell.font = Font(name="Calibri", size=10)
            if col_idx in [1, 2, 4, 5, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ans_cell = ws.cell(row=row_idx, column=8)
        ans_upper = str(effective_ans).upper()
        if "HIJAU" in ans_upper:
            ans_cell.fill = hijau_fill
            ans_cell.font = Font(name="Calibri", size=10, bold=True, color="155724")
        elif "TRANSISI" in ans_upper:
            ans_cell.fill = transisi_fill
            ans_cell.font = Font(name="Calibri", size=10, bold=True, color="856404")
        else:
            ans_cell.fill = tidak_fill
            ans_cell.font = Font(name="Calibri", size=10, bold=True, color="721C24")

    # Set column widths
    col_widths = {
        "A": 14, "B": 24, "C": 45, "D": 22, "E": 28, "F": 32,
        "G": 24, "H": 18, "I": 16, "J": 55, "K": 40, "L": 25,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 28
    for r_i in range(2, len(entries) + 2):
        ws.row_dimensions[r_i].height = 36

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
